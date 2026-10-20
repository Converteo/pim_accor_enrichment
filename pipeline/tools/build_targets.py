"""Generate prompts/enrichment/<family>/targets.json from the audit workbook (livrable 1.1) and the PIM model.

    python tools/build_targets.py "<...>/1.1 Rapprochement et enrichissements possibles - Audit Qualité PIM Hôtels.xlsx" \
        <...>/live_model.json --version enrich-v1

The workbook gives the 54 fields in scope with their priority (P1 site, P2 API, P3 not exposed); the PIM model
gives their type, list of values (CVL) and multi-value flag. Families, source columns and instructions are set here.
Re-run after a change of priorities, then review and commit the generated files.
"""

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SHEET = "2 - Champs enrichissables"

FAMILIES = {
    "general": [
        "ProductCheckIn", "ProductCheckOut", "ProductStaffLanguages", "ProductNonSmokingProperty", "ProductAirConditionInAllArea",
        "ProductCheckInCheckOutFacilities", "ProductFreeAdditionalServices", "ProductPayingAdditionalServices",
        "ProductFreeGeneralAmenities", "ProductFreeServices", "ProductTourismServicesDesks", "ProductFreeDeskEquipmentAmenities",
        "ProductFamilyAndChildrenServices", "ProductEntertainmentServices",
    ],
    "accessibility_safety": [
        "ProductReducedMobilityFacilities", "ProductWheelchairAccessible", "ProductHotelZoneAccessible", "ProductFeaturesAccessible",
        "ProductFireEmergencyInfo", "ProductStaffMonitoringInfo", "ProductAccessControlInfo", "ProductAlarmInfo",
    ],
    "dining": [
        "ProductFnBAvailableBar", "ProductServiceFullMenuInRoom", "ProductServicesAvailableKidsMenu",
        "BarType", "BreakfastCharged", "BreakfastCategory", "RestaurantTypeFood",
    ],
    "pool": [
        "ProductNbPools", "ProductAvailablePool", "PoolType", "PoolHeated", "PoolPaying", "PoolOverflowing", "PoolWater", "PoolFeatures",
    ],
    "parking": [
        "ParkingCharged", "ParkingTypeOfParking", "ParkingLocationOfParking", "ParkingCost", "ParkingPricingMode",
        "ParkingServices", "ParkingSecurityAccessibility", "ParkingAvailableChargingTerminal", "ParkingIndoorOutdoor",
    ],
    "wellness_fitness": [
        "ProductFitnessCenter", "ProductAvailableWellnessCenter", "WellnessCenterType", "WellnessCenterFreeWellnessCenterServices",
        "WellnessCenterChargedWellnessCenterServices", "FitnessCenterFreeWellServicesEquipments",
        "FitnessCenterChargedPersoLessonsSpecialCare", "SportType",
    ],
}

# Raw columns (snake_case) sent to the agent for each family
BOOKING_ALL = ["check_in", "check_out", "general", "internet", "food_drink", "wellness", "activities", "reception_services",
               "cleaning_services", "business_facilities", "entertainment_family_services", "safety_security", "accessibility",
               "languages_spoken", "services", "services_extras", "pets", "parking", "nb_piscines"]
POOLS = [f"piscine_{i}" for i in range(1, 16)]
SOURCES = {
    "general": {"booking_equipements": ["check_in", "check_out", "general", "internet", "reception_services", "cleaning_services",
                                              "business_facilities", "entertainment_family_services", "activities", "languages_spoken",
                                              "services", "services_extras", "pets"],
                      "expedia": ["policies", "important_info", "points_forts_sejour", "property_about"]},
    "accessibility_safety": {"booking_equipements": ["accessibility", "general", "safety_security"],
                               "expedia": ["important_info", "property_about"]},
    "dining": {"booking_equipements": ["food_drink", "general", "services"],
                     "expedia": ["important_info", "points_forts_sejour", "property_about"]},
    "pool": {"booking_equipements": ["nb_piscines", *POOLS, "wellness"], "booking_piscines": ["nb_pool"],
                "expedia": ["points_forts_sejour", "property_about", "important_info"]},
    "parking": {"booking_equipements": ["parking"], "expedia": ["policies", "important_info", "property_about"]},
    "wellness_fitness": {"booking_equipements": ["wellness", "activities"], "expedia": ["points_forts_sejour", "property_about", "important_info"]},
}

FALSE_IS_AN_ANSWER = {"ParkingCharged", "PoolPaying", "BreakfastCharged"}
HHMM = {"ProductCheckIn", "ProductCheckOut"}
INSTRUCTIONS = {
    "ProductCheckIn": "Heure de début du check-in (Booking « From 15:00 », Expedia « Check-in start time: 3:00 PM »).",
    "ProductCheckOut": "Heure limite du check-out (Booking « Until 12:00 », Expedia « Check-out before noon »).",
    "ProductStaffLanguages": "Langues parlées par le personnel ; n'ajoute que des langues de la liste PIM.",
    "ProductNbPools": "Nombre de piscines de l'hôtel (Booking NB_PISCINES ou fichier piscines NB_POOL).",
    "ProductAirConditionInAllArea": "True seulement si la source dit que la climatisation couvre tout l'hôtel ; « Air conditioning » seul ne suffit pas.",
    "ProductPayingAdditionalServices": "Services payants seulement si la source précise qu'ils sont payants (« additional charge »).",
    "ParkingCost": "Montant seul, sans devise ; signale la devise dans l'alerte.",
    "ParkingCharged": "False si le parking est gratuit (« Free parking »), True s'il est payant (« costs », « charge »).",
    "PoolType": "Intérieur ou extérieur, pour chaque piscine.",
    "BreakfastCharged": "False si le petit-déjeuner est inclus ou gratuit, True s'il est payant (« Fee for breakfast »).",
    "WellnessCenterChargedWellnessCenterServices": "Uniquement les services que la source dit payants.",
}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook")
    parser.add_argument("model_json")
    parser.add_argument("--version", default="enrich-v1")
    args = parser.parse_args()

    ws = load_workbook(args.workbook, read_only=True)[SHEET]
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    inventory = {r[header.index("Champ PIM")]: dict(zip(header, r)) for r in rows[1:] if r[header.index("Champ PIM")]}
    model = {f["fieldTypeId"]: (e["id"], f) for e in json.loads(Path(args.model_json).read_text()) for f in e["fieldTypes"]}

    missing = [f for fs in FAMILIES.values() for f in fs if f not in inventory or f not in model]
    if missing:
        raise SystemExit(f"Fields missing from the workbook or the model: {missing}")
    for family, family_fields in FAMILIES.items():
        fields = []
        for field in family_fields:
            entity, f = model[field]
            inv = inventory[field]
            fields.append({
                "pim_field": field, "entity": entity, "label": f["fieldTypeDisplayName"], "pim_type": f["fieldDataType"],
                "cvl_id": f.get("cvlId"), "multivalue": bool(f.get("isMultiValue")), "priority": inv["Priorité"],
                "format": "hh:mm" if field in HHMM else None, "false_is_an_answer": field in FALSE_IS_AN_ANSWER,
                "instruction": INSTRUCTIONS.get(field) or inv.get("Remarque") or None, "audit_source": inv.get("Source(s)"),
            })
        fields.sort(key=lambda c: (c["priority"], c["entity"] != "Product", c["pim_field"]))
        out = ROOT / "prompts" / "enrichment" / family / "targets.json"
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(json.dumps({"family": family, "prompt_version": args.version, "sources": SOURCES[family], "fields": fields},
                                  ensure_ascii=False, indent=2) + "\n")
        print(f"{out.relative_to(ROOT)}: {len(fields)} fields")


if __name__ == "__main__":
    main()
